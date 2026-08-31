#!/usr/bin/env python3
"""Build data/releases/*-draft.json from manually cleaned interim Excel files.

Source of truth: data/interim 2/ (user-curated). Falls back to data/interim 1/
when interim 2 only has artifacts (2024 country list, 2026 footnotes).

Requires: pip install openpyxl
Run from repo root: py scripts/build_releases_from_interim.py
"""

from __future__ import annotations

import ast
import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INTERIM2 = ROOT / "data" / "interim 2"
INTERIM1 = ROOT / "data" / "interim 1"
OUT_DIR = ROOT / "data" / "releases"

# interim 2 filename -> release id suffix (YYYY-draft)
INTERIM2_FILES: dict[str, str] = {
    "2008_aala_alpha-06-24-14.xlsx": "2008-draft",
    "2011_aala_alpha2.xlsx": "2011-draft",
    "2020_aala_alpha_1-26-21.xlsx": "2020-draft",
    "MY2021-AALA-Alphabetical_4-27-22.xlsx": "2021-draft",
    "MY2022-AALA-Alphabetical-4-11-24_0.xlsx": "2022-draft",
    "MY2023-AALA-Alphabetical-02042025.xlsx": "2023-draft",
    "MY2025-AALA-Alphabetical 4_7_2025.xlsx": "2025-draft",
}

# Years where interim 2 is not usable; use interim 1 instead.
INTERIM1_FALLBACK: dict[str, str] = {
    "2024-draft": "MY2024-AALA-Alphabetical-2.4.25.xlsx",
    "2026-draft": "MY2026-AALA-Alphabetical-1.29.26_0.xlsx",
}

COUNTRY_CODES: dict[str, str] = {
    "AU": "Australia",
    "AT": "Austria",
    "BE": "Belgium",
    "BR": "Brazil",
    "CH": "China",
    "CU": "Cuba",
    "CN": "Canada",
    "CZ": "Czech Republic",
    "DE": "Denmark",
    "F": "France",
    "FN": "Finland",
    "G": "Germany",
    "H": "Hungary",
    "I": "Italy",
    "ID": "Indonesia",
    "IN": "India",
    "J": "Japan",
    "K": "Korea",
    "M": "Mexico",
    "N": "Netherlands",
    "OT": "Other",
    "P": "Philippines",
    "PL": "Poland",
    "PO": "Portugal",
    "RO": "Romania",
    "RU": "Russia",
    "SI": "Singapore",
    "SL": "Slovakia",
    "SP": "Spain",
    "SW": "Sweden",
    "T": "Turkey",
    "TH": "Thailand",
    "UK": "Great Britain",
    "US": "United States",
    "USA": "United States",
}

BREAKDOWN_RE = re.compile(
    r"^\s*(?P<pct>\d+(?:\.\d+)?)\s*%\s*(?P<code>[A-Za-z]{1,3}|USA|UK|OT)\s*$"
)


def load_rows(path: Path) -> list[tuple]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def is_data_header(header: tuple) -> bool:
    if not header:
        return False
    first = str(header[0] or "").strip().lower()
    return first.startswith("manufacturer")


def parse_percent(value) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, str):
        s = value.strip().replace("%", "")
        if not s:
            return None
        try:
            num = float(s)
        except ValueError:
            return None
    elif isinstance(value, (int, float)):
        num = float(value)
    else:
        return None

    if 0 <= num <= 1:
        num *= 100
    return int(round(num))


def resolve_country(raw) -> str | None:
    if raw is None or raw == "":
        return None
    code = str(raw).strip().upper()
    if not code:
        return None
    if code in COUNTRY_CODES:
        return COUNTRY_CODES[code]
    # Already a full name in some older rows
    return str(raw).strip()


def parse_breakdown_cell(cell) -> dict | None:
    if cell is None or cell == "":
        return None
    text = str(cell).strip()
    m = BREAKDOWN_RE.match(text)
    if not m:
        return None
    country = resolve_country(m.group("code"))
    if not country:
        return None
    return {"country": country, "percent": int(round(float(m.group("pct"))))}


def origins(primary, secondary=None) -> list[dict] | None:
    out = []
    p = resolve_country(primary)
    if p:
        out.append({"country": p, "role": "primary"})
    s = resolve_country(secondary)
    if s:
        out.append({"country": s, "role": "secondary"})
    return out or None


def foreign_source_for_year(release_id: str) -> str:
    year = int(release_id.split("-")[0])
    if year >= 2026:
        return "major_sources_2026"
    if year >= 2011:
        return "major_sources_2026"
    return "derived"


def parse_modern_row(row: tuple, idx: int, release_id: str) -> dict | None:
    corporation = str(row[0] or "").strip()
    brand = str(row[1] or "").strip()
    model = str(row[2] or "").strip()
    if not corporation or not model:
        return None

    value = parse_percent(row[4])
    if value is None:
        return None

    breakdown = []
    for cell in (row[5], row[6]) if len(row) > 6 else ():
        item = parse_breakdown_cell(cell)
        if item:
            breakdown.append(item)

    has_breakdown = len(breakdown) > 0
    foreign_total = max(0, 100 - value)
    foreign_source = foreign_source_for_year(release_id)

    record: dict = {
        "id": str(idx),
        "region": "",
        "corporation": corporation,
        "brand": brand or corporation,
        "model": model,
        "value": value,
        "foreignPartsTotalPercent": foreign_total,
        "foreignPartsTotalSource": "derived",
        "foreignPartsSource": foreign_source if has_breakdown else "derived",
    }

    if has_breakdown:
        record["foreignPartsBreakdown"] = breakdown
    else:
        record["foreignPartsBreakdown"] = []
        record["foreignPartsNote"] = (
            "Country breakdown not in source; total derived from 100% − U.S./Canadian."
        )

    if len(row) > 7:
        fa = origins(row[7], row[8] if len(row) > 8 else None)
        if fa:
            record["finalAssemblyCountries"] = fa
    if len(row) > 9:
        eng = origins(row[9], row[10] if len(row) > 10 else None)
        if eng:
            record["engineOrigins"] = eng
    if len(row) > 11:
        trn = origins(row[11], row[12] if len(row) > 12 else None)
        if trn:
            record["transmissionOrigins"] = trn
    if len(row) > 3 and row[3]:
        record["vehicleTypePart567"] = str(row[3]).strip()

    return record


def parse_2008_row(row: tuple, idx: int, release_id: str) -> dict | None:
    corporation = str(row[0] or "").strip()
    brand = str(row[1] or "").strip()
    model = str(row[2] or "").strip()
    if not corporation or not model:
        return None

    value = parse_percent(row[4])
    if value is None:
        return None

    foreign_total = max(0, 100 - value)
    return {
        "id": str(idx),
        "region": "",
        "corporation": corporation,
        "brand": brand or corporation,
        "model": model,
        "value": value,
        "foreignPartsTotalPercent": foreign_total,
        "foreignPartsTotalSource": "derived",
        "foreignPartsSource": "derived",
        "foreignPartsBreakdown": [],
        "foreignPartsNote": (
            "Country breakdown not in source; total derived from 100% − U.S./Canadian."
        ),
    }


def convert_workbook(path: Path, release_id: str) -> list[dict]:
    rows = load_rows(path)
    if not rows:
        return []

    header = rows[0]
    if not is_data_header(header):
        raise ValueError(f"{path.name}: expected Manufacturers header, got {header[0]!r}")

    year = release_id.split("-")[0]
    parser = parse_2008_row if year == "2008" else parse_modern_row

    out: list[dict] = []
    idx = 1
    for row in rows[1:]:
        if not row or all(c is None or c == "" for c in row):
            continue
        record = parser(row, idx, release_id)
        if record:
            out.append(record)
            idx += 1
    return out


def write_release(release_id: str, rows: list[dict]) -> Path:
    out_path = OUT_DIR / f"{release_id}.json"
    out_path.write_text(json.dumps(rows, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return out_path


def main() -> None:
    built: list[tuple[str, Path, int, str]] = []

    for filename, release_id in INTERIM2_FILES.items():
        path = INTERIM2 / filename
        if not path.exists():
            print(f"skip missing {path}")
            continue
        rows = convert_workbook(path, release_id)
        out = write_release(release_id, rows)
        built.append((release_id, path, len(rows), "interim 2"))
        print(f"OK {release_id}: {len(rows)} rows <- interim 2/{filename}")

    for release_id, filename in INTERIM1_FALLBACK.items():
        path = INTERIM1 / filename
        if not path.exists():
            print(f"skip missing fallback {path}")
            continue
        rows = convert_workbook(path, release_id)
        out = write_release(release_id, rows)
        built.append((release_id, path, len(rows), "interim 1 (fallback)"))
        print(f"OK {release_id}: {len(rows)} rows <- interim 1/{filename}")

    print("\nBuilt releases:")
    for release_id, src, count, note in built:
        print(f"  {release_id}.json  ({count} rows)  [{note}]")


if __name__ == "__main__":
    main()
